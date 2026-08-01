# Copyright (c) 2025, NVIDIA CORPORATION.  All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Extract text content and convert deliverable files for reward computation.

Reads common document formats (.docx, .pdf, .xlsx, .pptx, .txt, etc.)
and returns their combined text so the LLM judge can score actual content
rather than just the agent's finish-tool summary.

Also provides PDF conversion via LibreOffice headless for visual judging
with multimodal models (e.g., Gemini 3 Pro).
"""

from __future__ import annotations

import base64
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any


MAX_TOTAL_CHARS = 20_000

# Caps on *text* content blocks in the multimodal judging path. ``read_deliverable_files``
# has always bounded its output at ``MAX_TOTAL_CHARS``; ``convert_deliverables_to_content_blocks``
# bounded nothing, so a single large text file went into the judge prompt whole. Observed:
# a 2.8 MB file produced a 720,898-token request against a 786,432-token judge, came back
# 400, and surfaced to the caller as a transient 500 -- so the task was not judged badly,
# it was not judged at all, and the run still exited 0.
#
# The largest genuine text deliverable in the GDPVal/Mercor corpus is 101 KB (per-task
# total also 101 KB), so these caps leave real submissions byte-for-byte intact and only
# bite pathological files. Truncating loses a tail; not truncating loses the task.
MAX_TEXT_BLOCK_CHARS = 200_000
MAX_TOTAL_TEXT_BLOCK_CHARS = 400_000
# Reserved per text block for its header and any truncation notice, so a directory
# of many small files cannot blow the request on framing alone.
TEXT_BLOCK_OVERHEAD_CHARS = 120

# Stirrup persists its own run state into the very directory it writes
# deliverables to: the finish payload, the full message history, and run
# metadata. ``.json`` is in ``TEXT_EXTS``, so without this filter the judge is
# handed the agent's own reasoning trace as part of the "submission" and can
# grade what the agent *said it did* rather than the artefact it produced -- a
# false positive for any "states/reports <value>" rubric criterion.
#
# This is the single definition: the agent that writes these files owns the
# list, and every judging path imports it from here (pairwise comparison in
# ``resources_servers/gdpval/comparison.py`` previously kept its own copy).
IGNORE_FILES = frozenset(
    {
        "finish_params.json",
        "history.json",
        "history.pkl",
        "metadata.json",
        "inprogress_history.json",
        "log.txt",
        "reference_files",
    }
)


def is_deliverable(path: Path) -> bool:
    """True if *path* is an agent-produced deliverable rather than run state."""
    return path.is_file() and path.name not in IGNORE_FILES


def read_deliverable_files(output_dir: str) -> str:
    """Read text from all deliverable files in *output_dir*.

    Returns a single string with sections like::

        === report.docx ===
        <extracted text>

        === data.xlsx ===
        <extracted text>

    Truncated to ~20k chars to stay within judge context limits.
    """
    output_path = Path(output_dir)
    if not output_path.is_dir():
        return ""

    files = sorted(p for p in output_path.iterdir() if p.name not in IGNORE_FILES)
    if not files:
        return ""

    parts: list[str] = []
    total_len = 0

    for fpath in files:
        if not fpath.is_file():
            continue

        ext = fpath.suffix.lower()
        try:
            text = _extract_text(fpath, ext)
        except Exception as exc:
            text = f"[Error reading {fpath.name}: {exc}]"

        if not text:
            continue

        section = f"=== {fpath.name} ===\n{text}"
        if total_len + len(section) > MAX_TOTAL_CHARS:
            remaining = MAX_TOTAL_CHARS - total_len
            if remaining > 200:
                section = section[:remaining] + "\n[...truncated]"
                parts.append(section)
            break

        parts.append(section)
        total_len += len(section)

    return "\n\n".join(parts)


def _extract_text(fpath: Path, ext: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    if ext in (".txt", ".md", ".csv", ".json", ".html", ".xml", ".log"):
        return _read_text(fpath)
    elif ext == ".docx":
        return _read_docx(fpath)
    elif ext == ".pdf":
        return _read_pdf(fpath)
    elif ext == ".xlsx":
        return _read_xlsx(fpath)
    elif ext == ".pptx":
        return _read_pptx(fpath)
    else:
        size = os.path.getsize(fpath)
        return f"[Binary file: {fpath.name}, {size} bytes]"


def _read_text(fpath: Path) -> str:
    return fpath.read_text(encoding="utf-8", errors="replace").strip()


def _read_docx(fpath: Path) -> str:
    from docx import Document

    doc = Document(str(fpath))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_pdf(fpath: Path) -> str:
    from pdfminer.high_level import extract_text

    return extract_text(str(fpath)).strip()


def _read_xlsx(fpath: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(fpath), read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(", ".join(cells))
        if rows:
            parts.append(f"Sheet: {sheet_name}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _read_pptx(fpath: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(fpath))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        texts.append(para.text)
        if texts:
            parts.append(f"Slide {i}:\n" + "\n".join(texts))
    return "\n\n".join(parts)


# ---------------------------------------------------------------------------
# PDF conversion for visual judging (Gemini 3 Pro)
# ---------------------------------------------------------------------------

OOXML_EXTS = {".docx", ".pptx", ".xlsx"}
# Legacy binary Office formats. LibreOffice renders these as readily as OOXML and
# ``resources_servers/gdpval/preconvert.py`` has always converted them (its
# ``LEGACY_OFFICE_EXTENSIONS``) -- but this set did not list them, so file_reader
# never looked for the sibling PDF preconvert had already written and the file
# reached no handler at all. Three such files in set A's c80a arm alone.
LEGACY_OFFICE_EXTS = {".doc", ".ppt", ".xls"}
OFFICE_EXTS = OOXML_EXTS | LEGACY_OFFICE_EXTS
# Extensions read as plain text. Deliberately generous: a file that reaches no
# handler in ``convert_deliverables_to_content_blocks`` emits NOTHING, and the
# judge reads that silence as non-delivery. Measured on set A (2,392 judged
# rollouts): 8 tasks whose deliverables were ``.ts`` / ``.ipynb`` / ``.toml`` /
# ``.zip`` averaged 0.1995 against 0.6401 for everything else, one of them graded
# "Missing: all 13 source modules" against a directory holding 31 files. Source
# and config files are the entire deliverable for a whole occupation category
# (Software Developers), so they are first-class here rather than an afterthought.
TEXT_EXTS = (
    {".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".html", ".htm", ".svg"}
    | {".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".properties", ".env"}
    | {".py", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".sql", ".r", ".tex", ".ipynb"}
    | {".ts", ".tsx", ".js", ".jsx", ".mjs", ".cjs", ".vue", ".css", ".scss", ".less"}
    | {".java", ".go", ".rs", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".kt", ".swift"}
    | {".rb", ".php", ".pl", ".lua", ".scala", ".jl", ".m", ".log", ".diff", ".patch"}
)
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".heic", ".heif"}
# Audio/video are only passed through when the judge is AV-capable (e.g. MiniMax
# M3); image-only judges can't decode them, so they are otherwise skipped.
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac"}
VIDEO_EXTS = {".mp4", ".mov", ".webm", ".mkv", ".avi"}

# Every extension with a dedicated branch in the dispatch loop below. Anything
# outside this set takes the unknown-extension path, which announces the file
# instead of dropping it. Keep in sync by construction, not by hand.
HANDLED_EXTS = TEXT_EXTS | OFFICE_EXTS | IMAGE_EXTS | AUDIO_EXTS | VIDEO_EXTS | {".pdf"}

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".heic": "image/heic",
    ".heif": "image/heif",
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".m4a": "audio/mp4",
    ".flac": "audio/flac",
    ".ogg": "audio/ogg",
    ".aac": "audio/aac",
    ".mp4": "video/mp4",
    ".mov": "video/quicktime",
    ".webm": "video/webm",
    ".mkv": "video/x-matroska",
    ".avi": "video/x-msvideo",
}

# Bytes probed when deciding whether an unknown-extension file is text. Enough to
# see a binary header without paying to read a large file twice.
TEXT_SNIFF_BYTES = 8192
# Cap on archive members listed. A manifest is context for grading "was the
# bundle delivered and does it contain X", not a deliverable in its own right.
MAX_ARCHIVE_ENTRIES = 200


def _human_size(nbytes: int) -> str:
    """``39466556`` -> ``37.6 MB (39,466,556 bytes)``; small sizes stay exact."""
    size = float(nbytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{nbytes:,} bytes" if unit == "B" else f"{size:.1f} {unit} ({nbytes:,} bytes)"
        size /= 1024
    return f"{nbytes:,} bytes"


def _sniffs_as_text(fpath: Path) -> bool:
    """True if an unknown-extension file looks like text we can hand the judge.

    An extension allowlist always lags reality -- a new framework, a ``Makefile``,
    a ``Dockerfile``, an extensionless script -- so unrecognised files are decided
    by content instead of by name. A NUL byte, or a meaningful density of
    undecodable bytes and control characters, means binary.
    """
    try:
        with fpath.open("rb") as fh:
            chunk = fh.read(TEXT_SNIFF_BYTES)
    except OSError:
        return False
    if not chunk or b"\x00" in chunk:
        return False
    text = chunk.decode("utf-8", errors="replace")
    # A replacement char appears legitimately where the probe cut a multi-byte
    # codepoint in half, so judge on density rather than on presence.
    noise = text.count("\ufffd") + sum(1 for ch in text if ord(ch) < 32 and ch not in "\t\n\r\f\v")
    return noise / len(text) < 0.01


def _archive_manifest(fpath: Path) -> str | None:
    """A listing of *fpath*'s members, or ``None`` if it is not a readable archive.

    A ``.zip`` deliverable is opaque as bytes, and a judge told only that it exists
    still cannot grade "the bundle contains X" -- which is what such rubrics ask.
    The member list is the part of its content the criteria actually reference.
    """
    import zipfile

    try:
        with zipfile.ZipFile(fpath) as zf:
            infos = zf.infolist()
    except Exception:
        # Not a zip, truncated, or encrypted. The caller falls back to announcing
        # the file by name and size, which is still better than silence.
        return None
    lines = [f"  {info.filename} ({info.file_size:,} bytes)" for info in infos[:MAX_ARCHIVE_ENTRIES]]
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        lines.append(f"  ... and {len(infos) - MAX_ARCHIVE_ENTRIES:,} more entries")
    return "\n".join(lines)


def _convert_office_to_pdf(fpath: Path, out_dir: Path | None = None) -> Path | None:
    """Convert a .docx/.xlsx/.pptx file to PDF using LibreOffice headless.

    Returns the path to the generated PDF, or None on failure.
    Uses a unique user profile to avoid lock conflicts in concurrent workers.

    With *out_dir*, the PDF is written there instead of next to *fpath*. Judging
    passes use that so reading a deliverables directory never also writes to it.

    Whitespace in the input filename makes LibreOffice's batch-convert mode
    silently drop the file (the URI it builds isn't percent-encoded), so we
    stage the input to a tempdir with a sanitized basename and move the PDF
    back to the intended location.
    """
    profile_dir = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    user_install = f"file://{profile_dir.as_posix()}"
    dest_dir = out_dir if out_dir is not None else fpath.parent
    out_pdf = dest_dir / (fpath.stem + ".pdf")
    stage_dir: Path | None = None
    input_path = fpath
    has_whitespace = any(c.isspace() for c in fpath.name)

    try:
        if has_whitespace:
            stage_dir = Path(tempfile.mkdtemp(prefix="lo-stage-"))
            safe_name = re.sub(r"\s+", "_", fpath.stem) + fpath.suffix
            input_path = stage_dir / safe_name
            shutil.copy2(fpath, input_path)
            lo_outdir = str(stage_dir)
        else:
            lo_outdir = str(dest_dir)

        cmd = [
            "libreoffice",
            "--headless",
            "--nologo",
            "--nolockcheck",
            "--nodefault",
            "--norestore",
            f"-env:UserInstallation={user_install}",
            "--convert-to",
            "pdf",
            "--outdir",
            lo_outdir,
            str(input_path),
        ]
        p = subprocess.run(cmd, check=False, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
        if stage_dir is not None:
            staged_pdf = stage_dir / (input_path.stem + ".pdf")
            if staged_pdf.exists():
                shutil.move(str(staged_pdf), str(out_pdf))
        if p.returncode != 0 or not out_pdf.exists():
            print(f"[file_reader] LibreOffice conversion failed for {fpath.name}: {p.stderr[:200]}", flush=True)
            return None
        return out_pdf
    except subprocess.TimeoutExpired:
        print(f"[file_reader] LibreOffice conversion timed out for {fpath.name}", flush=True)
        return None
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
        if stage_dir is not None:
            shutil.rmtree(stage_dir, ignore_errors=True)


def _pdf_bytes_to_image_text_blocks(
    pdf_bytes: bytes,
    *,
    render_dpi: int,
    max_pages: int,
    include_text: bool,
) -> list[dict[str, Any]] | None:
    """Rasterize PDF bytes to page-image + text blocks for image-only judges.

    Delegates to :mod:`resources_servers.gdpval.media_conversion` (imported
    lazily so this module stays importable even where that package isn't on the
    path). Returns ``None`` when the helper is unavailable or yields nothing, so
    the caller can fall back to the native ``application/pdf`` data URL.
    """
    try:
        from resources_servers.gdpval.media_conversion import pdf_bytes_to_blocks
    except ImportError:
        return None
    blocks = pdf_bytes_to_blocks(
        pdf_bytes,
        dpi=render_dpi,
        max_pages=max_pages,
        include_text=include_text,
    )
    return blocks or None


def _av_block(mime: str, data: bytes, *, ext: str, file_type: str, openai_native: bool) -> dict[str, Any]:
    """Build an audio/video content block in the judge's dialect.

    Delegates to :mod:`resources_servers.gdpval.media_conversion` (imported
    lazily so this module stays importable where that package isn't on the
    path). Falls back to the legacy ``image_url`` data URL if the helper is
    unavailable — correct for frontier judges, and a benign degradation for
    self-hosted ones (which is only reachable when the gdpval package is
    present anyway).
    """
    try:
        from resources_servers.gdpval.media_conversion import audio_video_block

        return audio_video_block(mime, data, ext=ext, file_type=file_type, openai_native=openai_native)
    except ImportError:
        b64 = base64.b64encode(data).decode("ascii")
        return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}


def _fair_text_allowances(sizes: list[int], total_budget: int, per_file_cap: int) -> list[int]:
    """Max-min fair split of *total_budget* across text deliverables of *sizes*.

    Smallest first: each file may take an equal share of what is left divided by how
    many files still need one. Files under their share take only what they need and
    hand the surplus back, so a small real deliverable is never starved by a large
    one that happens to sort earlier. Returned in the order *sizes* was given.
    """
    allowances = [0] * len(sizes)
    remaining = total_budget
    order = sorted(range(len(sizes)), key=lambda i: sizes[i])
    for position, idx in enumerate(order):
        still_to_serve = len(order) - position
        share = remaining // still_to_serve
        take = max(0, min(sizes[idx], per_file_cap, share))
        allowances[idx] = take
        remaining -= take
    return allowances


def convert_deliverables_to_content_blocks(
    output_dir: str,
    *,
    media_mode: str = "native_pdf",
    render_dpi: int = 150,
    max_pages: int = 50,
    include_text: bool = True,
    audio_capable: bool = False,
    video_capable: bool = False,
) -> list[dict[str, Any]]:
    """Convert deliverable files to OpenAI-compatible content blocks for multimodal judging.

    Returns a list of content blocks suitable for the ``content`` field of an
    OpenAI chat message. Each file becomes either:

    - A text block (for .txt/.md/.csv etc.)
    - An image_url block with base64 data URL (for PDFs, images, converted Office docs)

    Office documents (.docx/.xlsx/.pptx) are first converted to PDF via
    LibreOffice headless so the judge sees rendered formatting/tables/charts.

    *media_mode* selects how PDFs/Office docs are presented: ``"native_pdf"``
    (default) sends them as ``application/pdf`` data URLs for frontier judges;
    ``"images_and_text"`` rasterizes each page to a PNG block and attaches the
    extracted text, for image-only local VLM judges (e.g. a gym-spawned Kimi
    K2.6). *render_dpi*, *max_pages*, and *include_text* tune that rendering.

    *audio_capable* / *video_capable* forward audio / video files (respectively)
    to judges that read that modality — tracked SEPARATELY because MiniMax-M3
    reads video but not audio (a video-only judge keeps video, skips audio). An
    unreadable modality's file is skipped. The AV block dialect follows
    *media_mode*: ``native_pdf`` (frontier judges, e.g. Gemini) uses an
    ``image_url`` data URL, while ``images_and_text`` (self-hosted vLLM judges)
    uses the standard ``video_url`` / ``input_audio`` content types vLLM routes to
    the media tower.
    """
    output_path = Path(output_dir)
    if not output_path.is_dir():
        return []

    images_and_text = media_mode == "images_and_text"

    blocks: list[dict[str, Any]] = []
    scratch_dirs: list[Path] = []  # tempdirs holding PDFs this pass converted

    entries = sorted(output_path.iterdir())

    # ``Report.docx`` and ``Report.pptx`` both map to ``Report.pdf``. A sibling PDF
    # is then ambiguous -- it can only be the render of one of them -- so reusing it
    # for the other would show the judge the first file's content twice and the
    # second file's content never. Count the stems and only take the fast path when
    # the mapping is unambiguous.
    office_stem_counts: dict[str, int] = {}
    for entry in entries:
        if is_deliverable(entry) and entry.suffix.lower() in OFFICE_EXTS:
            office_stem_counts[entry.stem] = office_stem_counts.get(entry.stem, 0) + 1

    # A sidecar render (``Plan.pptx.pdf``) is presented as part of its source
    # Office file below, so it must not also be emitted as a standalone PDF --
    # that would show the judge the same slides twice and invite double-counting
    # against the rubric.
    sidecar_pdfs = {
        entry.with_name(entry.name + ".pdf")
        for entry in entries
        if is_deliverable(entry) and entry.suffix.lower() in OFFICE_EXTS
    }

    # Allot the aggregate budget up front rather than first-come-first-served. Files
    # are visited in sorted order, so a spend-as-you-go budget would let two big
    # ``000_*.txt`` files starve a real ``Report.md`` -- renaming a file would change
    # the score. Max-min fair shares make a deliverable's allowance independent of
    # what it is called. Per-block overhead (header, truncation notice) is reserved
    # before allotting, because it is tokens too.
    # Files handed to the judge as text: a recognised text extension, or an
    # unrecognised one whose content sniffs as text. Decided ONCE here so the
    # allotment below and the dispatch loop cannot disagree -- a file emitted as
    # text but missing from this list would be allotted nothing and truncated to
    # zero characters, which is the same silence the else-branch exists to end.
    texty = {
        p
        for p in entries
        if is_deliverable(p)
        and p not in sidecar_pdfs
        and (p.suffix.lower() in TEXT_EXTS or (p.suffix.lower() not in HANDLED_EXTS and _sniffs_as_text(p)))
    }

    text_files = [p for p in entries if p in texty]
    body_budget = max(0, MAX_TOTAL_TEXT_BLOCK_CHARS - len(text_files) * TEXT_BLOCK_OVERHEAD_CHARS)
    allowance_by_name = dict(
        zip(
            (p.name for p in text_files),
            _fair_text_allowances([p.stat().st_size for p in text_files], body_budget, MAX_TEXT_BLOCK_CHARS),
        )
    )

    emitted_chars = 0
    omitted_names: list[str] = []

    def _text_block(header: str, body: str, allowance: int) -> dict[str, Any] | None:
        """A text block for *body*, trimmed to *allowance*.

        Returns ``None`` once the aggregate budget is spent: with enough files even
        the headers alone would blow the request, so past that point they are counted
        into a single summary block rather than named one by one.
        """
        nonlocal emitted_chars
        if emitted_chars >= MAX_TOTAL_TEXT_BLOCK_CHARS:
            return None
        cap = max(0, min(MAX_TEXT_BLOCK_CHARS, allowance))
        if len(body) > cap:
            shown = body[:cap] + f"\n[...truncated: {len(body) - cap:,} of {len(body):,} chars omitted]"
        else:
            shown = body
        block = {"type": "text", "text": f"\n{header}\n{shown}"}
        emitted_chars += len(block["text"])
        return block

    for fpath in entries:
        if not is_deliverable(fpath) or fpath in sidecar_pdfs:
            continue

        ext = fpath.suffix.lower()

        try:
            if fpath in texty:
                text = fpath.read_text(encoding="utf-8", errors="replace").strip()
                if text:
                    block = _text_block(f"{fpath.name}:", text, allowance_by_name.get(fpath.name, 0))
                    blocks.append(block) if block else omitted_names.append(fpath.name)
                else:
                    # An empty file is a real outcome and the judge should grade it
                    # as one -- but it must be told, or the file is indistinguishable
                    # from one that was never written.
                    blocks.append({"type": "text", "text": f"\n{fpath.name}: [present but EMPTY (0 bytes of content)]"})

            elif ext in OFFICE_EXTS:
                # A preconversion pass may already have written a sibling PDF
                # (see resources_servers/gdpval/preconvert.py, which guards on
                # exactly this). Reuse it when it unambiguously belongs to this
                # file, rather than regenerating over the top of a corpus we did
                # not create -- a judging pass that rewrites its own inputs makes
                # the next run over the same tree a different experiment.
                # Preconvert writes same-stem Office files to an injective
                # sidecar (``Plan.pptx`` -> ``Plan.pptx.pdf``) precisely because
                # ``Plan.pdf`` cannot say which file it renders. Prefer it: it is
                # unambiguous by construction, whatever the stem counts are.
                sidecar = fpath.with_name(fpath.name + ".pdf")
                existing_pdf = fpath.with_suffix(".pdf")
                if sidecar.is_file():
                    pdf_path = sidecar
                elif existing_pdf.is_file() and office_stem_counts.get(fpath.stem, 0) == 1:
                    pdf_path = existing_pdf
                else:
                    # Convert into a tempdir, never next to the original: that keeps
                    # the deliverables directory read-only for judging and stops two
                    # same-stem Office files fighting over one output name.
                    scratch = Path(tempfile.mkdtemp(prefix="deliverable-pdf-"))
                    scratch_dirs.append(scratch)
                    pdf_path = _convert_office_to_pdf(fpath, out_dir=scratch)
                if pdf_path and pdf_path.exists():
                    data = pdf_path.read_bytes()
                    if images_and_text:
                        rendered = _pdf_bytes_to_image_text_blocks(
                            data, render_dpi=render_dpi, max_pages=max_pages, include_text=include_text
                        )
                        if rendered is not None:
                            blocks.append({"type": "text", "text": f"\n{fpath.name} (rendered from PDF):"})
                            blocks.extend(rendered)
                            continue
                    b64 = base64.b64encode(data).decode("ascii")
                    blocks.append({"type": "text", "text": f"\n{fpath.name} (converted to PDF):"})
                    blocks.append(
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:application/pdf;base64,{b64}"},
                        }
                    )
                else:
                    # Fallback to text extraction
                    # No pre-allotted share: an Office file only lands here when its
                    # conversion failed, so it competes for whatever budget is left.
                    text = _extract_text(fpath, ext)
                    if text:
                        block = _text_block(f"{fpath.name} (text fallback):", text, MAX_TEXT_BLOCK_CHARS)
                        blocks.append(block) if block else omitted_names.append(fpath.name)

            elif ext == ".pdf":
                data = fpath.read_bytes()
                if images_and_text:
                    rendered = _pdf_bytes_to_image_text_blocks(
                        data, render_dpi=render_dpi, max_pages=max_pages, include_text=include_text
                    )
                    if rendered is not None:
                        blocks.append({"type": "text", "text": f"\n{fpath.name}:"})
                        blocks.extend(rendered)
                        continue
                b64 = base64.b64encode(data).decode("ascii")
                blocks.append({"type": "text", "text": f"\n{fpath.name}:"})
                blocks.append(
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:application/pdf;base64,{b64}"},
                    }
                )

            elif ext in IMAGE_EXTS:
                mime = MIME_TYPES.get(ext, "image/png")
                data = fpath.read_bytes()
                b64 = base64.b64encode(data).decode("ascii")
                blocks.append({"type": "text", "text": f"\n{fpath.name}:"})
                blocks.append(
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{b64}"},
                    }
                )

            elif ext in AUDIO_EXTS or ext in VIDEO_EXTS:
                # Forward AV only to judges that read that SPECIFIC modality, gated
                # independently: video needs *video_capable*, audio *audio_capable*
                # — so MiniMax-M3 (video yes, audio no) keeps video and skips audio.
                # Among frontier judges only Gemini reads AV. Gemini (native_pdf)
                # accepts AV as an image_url data URL; self-hosted vLLM judges
                # (images_and_text) need the standard video_url / input_audio types,
                # which vLLM routes to the media tower.
                is_video = ext in VIDEO_EXTS
                file_type = "VIDEO" if is_video else "AUDIO"
                if video_capable if is_video else audio_capable:
                    mime = MIME_TYPES.get(ext, "application/octet-stream")
                    data = fpath.read_bytes()
                    blocks.append({"type": "text", "text": f"\n{fpath.name}:"})
                    blocks.append(_av_block(mime, data, ext=ext, file_type=file_type, openai_native=images_and_text))
                else:
                    # The gate is a capability limit, not an absence. Emitting
                    # nothing made the judge assert the file was never produced: on
                    # set A the five .wav tasks averaged 0.213 against 0.630 for
                    # everything else, one graded "there is no actual .wav file
                    # included" while a 39 MB .wav sat in the directory. Say the
                    # file is there, and say why it cannot be played.
                    verb = "watching" if is_video else "listening"
                    blocks.append(
                        {
                            "type": "text",
                            "text": (
                                f"\n{fpath.name}: [{file_type} deliverable, {_human_size(fpath.stat().st_size)} "
                                f"— present on disk, but this judge cannot decode {file_type.lower()}. "
                                f"Do NOT treat it as missing or unproduced. Grade the criteria that do not "
                                f"require {verb}; mark the rest unverifiable rather than unmet.]"
                            ),
                        }
                    )

            else:
                # No handler for this extension. Before this branch existed the loop
                # emitted nothing at all and the judge reported non-delivery -- "The
                # required Northstar_DeltaFix_Bundle.zip is missing entirely" against
                # an 8,357-byte file on disk. Whatever the allowlist above holds, an
                # unrecognised deliverable must still be announced; silence must
                # never be readable as absence.
                size = fpath.stat().st_size
                if size == 0:
                    body = "[present but EMPTY (0 bytes)]"
                else:
                    manifest = _archive_manifest(fpath)
                    if manifest is not None:
                        body = f"[archive, {_human_size(size)} — present, NOT missing; contents:]\n{manifest}"
                    else:
                        body = f"[deliverable present, {_human_size(size)} — NOT missing; content not extractable in this format]"
                block = _text_block(f"{fpath.name}:", body, MAX_TEXT_BLOCK_CHARS)
                blocks.append(block) if block else omitted_names.append(fpath.name)
        except Exception as exc:
            blocks.append({"type": "text", "text": f"\n{fpath.name}: [Error: {exc}]"})

    if omitted_names:
        shown = ", ".join(omitted_names[:20])
        if len(omitted_names) > 20:
            shown += f", and {len(omitted_names) - 20} more"
        blocks.append(
            {
                "type": "text",
                "text": f"\n[{len(omitted_names)} text deliverable(s) omitted, budget exhausted: {shown}]",
            }
        )

    # Clean up only what this pass created, which lives entirely in tempdirs.
    for scratch in scratch_dirs:
        shutil.rmtree(scratch, ignore_errors=True)

    return blocks
